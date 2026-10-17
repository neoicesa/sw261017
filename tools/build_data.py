#!/usr/bin/env python3
"""자율성과지표 엑셀 → site/data/data.js 변환 도구

사용법:
    pip install openpyxl
    python tools/build_data.py "자율성과지표_설정_및_목표_달성계획.xlsx"

전제(2026.10.12. 엑셀과 같은 구조):
  - 시트: 총괄표, ①~④ 지수 시트(시트 이름 순서대로 2~5번째)
  - 지수 시트 2행이 머리글: 구분 / 하위지표·하위항목 / 가중치 / 기준값 / 1차년도 목표 / 1차년도 실적 / 2차년도 목표 / … / 3차년도 목표 / 단위·비고
  - A열 구분값: 지수 / 하위지표 / 하위항목
대시보드는 현재 1차년도(’25) 목표·실적을 읽습니다.
"""
import csv
import io
import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit('openpyxl이 필요합니다:  pip install openpyxl')

# 하위지표 산출식(엑셀 수식을 풀어 쓴 설명). 구조가 바뀌면 빈칸으로 둡니다.
FORMULA = {
    '①': {'(1)': 'MoA-X 트랙 이수자 수 ÷ 참여학과 재학생 수 × 100', '(2)': 'AI 진단지 적용 수 ÷ AI 적용 진단지 총 수 × 100', '(3)': '직접 입력(%)', '(4)': '0.5×JA 교원 확보율 + 0.5×유연학사 운영률'},
    '②': {'(1)': 'AID 교양교과목 개발·개편 건수 ÷ 교양 교과목 전체 개수 × 100', '(2)': '0.3×핵심역량 진단 점수 + 0.3×AID 재학생 수료율 + 0.2×디지털 배지 취득률 + 0.2×AID 교직원 연수 이수율', '(3)': '0.5×AI 교수법 참여 교원 비율 + 0.5×하이브리드·AI 콘텐츠 적용 교과목 비율', '(4)': '0.5×평생지도교수제 참여율 + 0.5×온라인·하이브리드 상담 증가율'},
    '③': {'(1)': '0.4×현장실습 참여율 + 0.3×캡스톤디자인 참여율 + 0.3×주문식 교육과정 운영 학과 비율', '(2)': '현장실습 참여학생 취업자 수 ÷ 현장실습 참여자 수 × 100', '(3)': '0.3×디지털 교육 참여율 + 0.3×디지털 교육 수료율 + 0.4×교육 만족도 점수', '(4)': '[0.5×(가족회사 누적 수 ÷ ’27 목표) + 0.5×(신규 MOU 건수 ÷ ’27 목표)] × 100'},
    '④': {'(1)': '[0.4×(유학생 ÷ ’27 목표) + 0.3×(TOPIK ÷ ’27 목표) + 0.3×(해외 MOU ÷ ’27 목표)] × 100', '(2)': '0.5×멘토링 교원 참여율 + 0.5×환경보호 캠페인 참여율', '(3)': '0.4×공동학점 운영 증가율 + 0.3×공동학점 학생 참여율 + 0.3×OCW 달성률', '(4)': '직접 입력(%)'},
}
# 비율형 하위지표의 분모 항목(엑셀 수식 기준): (지수, 하위지표) → 분모 항목 기호
DENOMINATOR = {('①', '(1)'): '나.', ('①', '(2)'): '나.', ('②', '(1)'): '나.', ('③', '(2)'): '가.'}


def v(x):
    return '' if x is None else x


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    kpi, items, indices = [], [], []
    for ws in wb.worksheets[1:5]:
        hdr = [c.value for c in ws[2]]
        col = lambda pre: [i for i, h in enumerate(hdr) if h and str(h).startswith(pre)][0]
        iB, iT1, iA1, iT2, iT3, iN = col('기준값'), col('1차년도 목표'), col('1차년도 실적'), col('2차년도 목표'), col('3차년도 목표'), col('단위')
        rows = list(ws.iter_rows(min_row=3, values_only=True))
        name = rows[0][1].strip()
        mark = name[0]
        sub = None
        for r in rows:
            if r[0] == '하위지표':
                sub = r[1].strip()
                key = sub.split(' ')[0]
                kpi.append({'indicator': sub, 'index': name, 'index_no': mark, 'weight': v(r[2]), 'baseline': v(r[iB]), 'target': v(r[iT1]), 'actual': v(r[iA1]),
                            'target_y2': v(r[iT2]), 'target_y3': v(r[iT3]), 'definition': FORMULA.get(mark, {}).get(key, ''), 'note': (r[iN] or '').strip()})
            elif r[0] == '하위항목' and sub:
                item = r[1].strip()
                k = (mark, sub.split(' ')[0])
                role = ('분모' if item.startswith(DENOMINATOR[k]) else '분자') if k in DENOMINATOR else ('가중' if r[2] is not None else '입력')
                items.append({'index_no': mark, 'indicator': sub, 'item': item, 'weight': v(r[2]), 'baseline': v(r[iB]), 'target': v(r[iT1]), 'actual': v(r[iA1]),
                              'target_y2': v(r[iT2]), 'target_y3': v(r[iT3]), 'note': (r[iN] or '').strip(), 'role': role})
    t = wb.worksheets[0]
    for tr in t.iter_rows(min_row=3, values_only=True):
        if tr[0] == '자율성과지표':
            indices.append({'index': tr[1], 'baseline': v(tr[3]), 'target_y1': v(tr[4]), 'actual_y1': v(tr[5]), 'target_y2': v(tr[7]), 'target_y3': v(tr[8]),
                            'orig_target_y2': v(tr[9]), 'orig_target_y3': v(tr[10])})
    # 대시보드는 CSV에서 읽은 것과 같은 문자열 값을 씁니다
    as_text = lambda rows: [{k: ('' if x == '' else str(x)) for k, x in r.items()} for r in rows]
    data = {'kpi': as_text(kpi), 'indices': as_text(indices), 'items': as_text(items), 'budget': [], 'programs': [], 'participation': []}
    out = Path(__file__).resolve().parent.parent / 'site' / 'data' / 'data.js'
    out.write_text('/* AI사업단 성과관리 대시보드 자료 — 원본: ' + Path(path).name + ' 1차년도(’25)\n   tools/build_data.py 로 다시 만들 수 있습니다. */\nwindow.DASH_DATA = ' + json.dumps(data, ensure_ascii=False, indent=1) + ';\n', encoding='utf-8')
    print(f'저장: {out}  (지수 {len(indices)}개, 하위지표 {len(kpi)}개, 하위항목 {len(items)}개)')


if __name__ == '__main__':
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
